"""Seed the public PDF/PPTX export scenario (follow-up to seed_wp7.py).

Run from backend/ against the running stack's DB, after an admin user and
org exist (tools/agent/seed_org.py). Owner = that admin. All reports are
artifact_visibility=public. Prints a JSON of ids.

  identity   page dashboard, one identity-scoped query, rows SECRET-ROW-*
  plain      page dashboard, one plain query, rows PUBLIC-ROW-*
  mixed      page dashboard, identity query (SECRET-ROW-*) + plain (PUBLIC-ROW-*)
  deck_id    slides artifact over the identity query; the .pptx on disk bakes
             the SECRET-ROW-* rows, the way slide generation bakes the
             owner's snapshot into the deck
  deck_plain slides artifact over a plain query (PUBLIC-ROW-*)
"""
import asyncio
import json
import sys
import uuid

sys.path.insert(0, ".")
import main  # noqa: F401,E402
from sqlalchemy import select  # noqa: E402

from app.dependencies import async_session_maker  # noqa: E402
from app.models.organization import Organization  # noqa: E402
from app.models.query import Query  # noqa: E402
from app.models.report import Report  # noqa: E402
from app.models.step import Step  # noqa: E402
from app.models.user import User  # noqa: E402
from app.models.visualization import Visualization  # noqa: E402
from app.models.widget import Widget  # noqa: E402
from app.services.report_pdf_service import ReportPdfService  # noqa: E402
from tests.fixtures.artifact import seed_artifact  # noqa: E402

IDENTITY = [{"name": "email", "type": "string", "source": "identity"}]

# Renders every visualization as a table + bar chart, so the PDF text layer
# carries each row label that reached the page.
CODE = """<script type="text/babel">
function App() {
  const data = useArtifactData();
  if (!data) return <div>loading</div>;
  return (<div className="p-6"><h1 className="text-2xl font-bold">{data.report.title}</h1>
    {data.visualizations.map((v, k) => {
      const rows = v.rows || [];
      return (<div key={k}><h2>{v.title}: {rows.length} rows</h2>
        <EChart option={{xAxis:{type:'category',data:rows.map(r=>r.name)},yAxis:{type:'value'},
          series:[{type:'bar',data:rows.map(r=>r.amount)}]}} style={{height:200}} />
        <table><tbody>{rows.map((r,i)=><tr key={i}><td>{r.name}</td><td>{r.amount}</td></tr>)}</tbody></table>
      </div>);
    })}
    <p>END-OF-DASHBOARD</p></div>);
}
ReactDOM.createRoot(document.getElementById('root')).render(<App />);
</script>"""


async def query(db, org, user, report, title, params, prefix):
    s = uuid.uuid4().hex[:6]
    w = Widget(title=title, slug=f"w-{s}", report_id=report.id)
    db.add(w); await db.flush()
    st = Step(title=title, slug=f"s-{s}", status="success", widget_id=w.id,
              data={"rows": [{"name": f"{prefix}-{i}", "amount": 10 * (i + 1)} for i in range(4)],
                    "columns": [{"field": "name"}, {"field": "amount"}]},
              data_model={"type": "bar_chart"}, applied_params={"email": user.email} if params else None)
    db.add(st); await db.flush()
    q = Query(title=title, report_id=report.id, widget_id=w.id, default_step_id=st.id,
              parameters=params, organization_id=org.id)
    db.add(q); await db.flush()
    st.query_id = q.id
    v = Visualization(title=title, status="success", report_id=report.id, query_id=q.id,
                      view={"type": "bar_chart"})
    db.add(v); await db.flush()
    return q, v


async def report(db, org, user, title, queries, mode="page"):
    r = Report(title=title, slug=f"wp7x-{uuid.uuid4().hex[:6]}", organization_id=org.id,
               user_id=user.id, status="published", artifact_visibility="public")
    db.add(r); await db.flush()
    made = [await query(db, org, user, r, t, p, pre) for t, p, pre in queries]
    content = {"code": CODE if mode == "page" else "prs = Presentation()",
               "visualization_ids": [str(v.id) for _, v in made]}
    a = await seed_artifact(db, report_id=r.id, user_id=user.id, organization_id=org.id,
                            mode=mode, title=title, content=content)
    if mode == "slides":
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4)).text_frame
        for _, prefix_rows in [(t, pre) for t, _, pre in queries]:
            for i in range(4):
                box.add_paragraph().text = f"{prefix_rows}-{i}  {10 * (i + 1)}"
        ReportPdfService.PPTX_DIR.mkdir(parents=True, exist_ok=True)
        path = ReportPdfService.PPTX_DIR / f"{a.id}.pptx"
        prs.save(path)
        a.pptx_path = str(path)
    return {"report_id": str(r.id), "artifact_id": str(a.id),
            "query_ids": [str(q.id) for q, _ in made]}


async def main_():
    async with async_session_maker() as db:
        user = (await db.execute(select(User).where(User.email == "admin@example.com"))).scalar_one()
        org = (await db.execute(select(Organization))).scalars().first()
        out = {
            "identity": await report(db, org, user, "Identity dash",
                                     [("Identity sales", IDENTITY, "SECRET-ROW")]),
            "plain": await report(db, org, user, "Plain dash",
                                  [("Plain sales", [], "PUBLIC-ROW")]),
            "mixed": await report(db, org, user, "Mixed dash",
                                  [("Identity sales", IDENTITY, "SECRET-ROW"),
                                   ("Plain sales", [], "PUBLIC-ROW")]),
            "deck_id": await report(db, org, user, "Identity deck",
                                    [("Identity sales", IDENTITY, "SECRET-ROW")], mode="slides"),
            "deck_plain": await report(db, org, user, "Plain deck",
                                       [("Plain sales", [], "PUBLIC-ROW")], mode="slides"),
        }
        await db.commit()
        print(json.dumps(out))


asyncio.run(main_())
