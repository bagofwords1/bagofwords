"""Probe the public PDF/PPTX exports seeded by seed_wp7_exports.py.

Usage (backend venv): python probe_exports.py ids.json BASE_URL OWNER OUT_DIR
(OWNER: a JWT, or the JSON tools/agent/seed_org.py printed)

For every seeded report, downloads /r/{id}/export_pdf and /r/{id}/export_pptx
anonymously and as the owner, extracts the file's text (pypdfium2 /
python-pptx) and counts SECRET-ROW-* (identity-scoped, owner only) and
PUBLIC-ROW-* (plain shared) row labels. Also prints what the public JSON view
serves the anonymous reader, for parity.
"""
import json
import re
import sys
from pathlib import Path

import httpx

ids = json.load(open(sys.argv[1]))
base, token, out = sys.argv[2].rstrip("/"), sys.argv[3], Path(sys.argv[4])
out.mkdir(parents=True, exist_ok=True)
if token.endswith(".json"):
    token = json.load(open(token))["admin"]["token"]


def text_of(kind: str, path: Path) -> str:
    if kind == "pdf":
        import pypdfium2
        doc = pypdfium2.PdfDocument(str(path))
        return "\n".join(p.get_textpage().get_text_range() for p in doc)
    from pptx import Presentation
    return "\n".join(sh.text_frame.text for s in Presentation(str(path)).slides
                     for sh in s.shapes if sh.has_text_frame)


def labels(text: str, prefix: str) -> int:
    return len(set(re.findall(prefix + r"-\d", text)))


with httpx.Client(base_url=base, timeout=360) as c:
    for name, rec in ids.items():
        rid = rec["report_id"]
        json_rows = []
        for qid in rec["query_ids"]:
            d = c.get(f"/api/r/{rid}/queries/{qid}/step").json()
            json_rows.append(len((d.get("data") or {}).get("rows", [])))
        print(f"== {name}: public JSON view rows per query (anonymous) = {json_rows}")
        for kind in ("pdf", "pptx"):
            for who, headers in (("anonymous", {}), ("owner", {"Authorization": f"Bearer {token}"})):
                r = c.get(f"/api/r/{rid}/export_{kind}", headers=headers)
                if r.status_code != 200:
                    print(f"   {kind:4} {who:9} http {r.status_code}  {r.text[:90]}")
                    continue
                f = out / f"{name}-{who}.{kind}"
                f.write_bytes(r.content)
                t = text_of(kind, f)
                print(f"   {kind:4} {who:9} http 200  SECRET-ROW labels={labels(t, 'SECRET-ROW')} "
                      f"PUBLIC-ROW labels={labels(t, 'PUBLIC-ROW')}")
