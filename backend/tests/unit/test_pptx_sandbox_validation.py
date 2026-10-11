"""The slides (python-pptx) sandbox must be no softer than the main Python
sandbox: any private/dunder attribute access and any hardcoded-path file read
the main validator rejects is rejected before the deck code ever runs, while
ordinary python-pptx code keeps working."""

import uuid

import pytest

from app.ai.code_execution.code_execution import UnsafePythonError, validate_python_code
from app.ai.code_execution.pptx_executor import PptxCodeExecutor, validate_pptx_code

DECK_PREAMBLE = """
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
"""

# Attribute names outside the enumerated FORBIDDEN_ATTRIBUTES list — the
# wholesale private/dunder rule is what must catch them.
UNLISTED_PRIVATE_ATTRS = [
    "__reduce__",
    "__reduce_ex__",
    "__base__",
    "__getattribute__",
    "__init_subclass__",
    "__subclasshook__",
    "__traceback__",
    "_element",
    "_spTree",
]

HARDCODED_READS = [
    'import pandas as pd\ndf = pd.read_csv("/etc/hostname")',
    'import numpy as np\na = np.loadtxt(f"/tmp/{1}")',
]


def _run(code, out_path):
    return PptxCodeExecutor().execute_pptx_code(
        code=code,
        visualizations=[],
        report={"id": "r", "title": "T", "theme": None},
        output_path=out_path,
    )


@pytest.mark.parametrize("attr", UNLISTED_PRIVATE_ATTRS)
def test_private_attribute_access_rejected_like_main_sandbox(attr):
    code = DECK_PREAMBLE + f"x = slide.{attr}\n"
    with pytest.raises(UnsafePythonError):
        validate_python_code(code)
    with pytest.raises(UnsafePythonError):
        validate_pptx_code(code)


@pytest.mark.parametrize("code", HARDCODED_READS)
def test_hardcoded_path_reads_rejected_like_main_sandbox(code):
    with pytest.raises(UnsafePythonError):
        validate_python_code(code)
    with pytest.raises(UnsafePythonError):
        validate_pptx_code(code)


def test_rejected_deck_code_never_executes(tmp_path):
    out = tmp_path / f"{uuid.uuid4().hex}.pptx"
    code = DECK_PREAMBLE + "x = slide.__reduce__\nprs.save(_pptx_output_path)\n"
    with pytest.raises(UnsafePythonError):
        _run(code, out)
    assert not out.exists()


# Every pattern the slides prompt's quick reference teaches (create_artifact.
# _build_slides_prompt), in one deck. If the prompt ever teaches something the
# validator rejects, every generated deck pays a repair round-trip — this fails
# first.
PROMPT_REFERENCE_DECK = """
from pptx.enum.dml import MSO_THEME_COLOR
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])
fill = slide.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(15, 23, 42)
pic = slide.shapes.add_picture(image(image_ids[0]), Inches(0), Inches(0), width=Inches(13.333))
band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(5), Inches(13.333), Inches(2.5))
band.fill.solid()
band.fill.fore_color.rgb = RGBColor(0, 31, 63)
band.line.fill.background()
tf = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11), Inches(1)).text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = report.get("title") or "Deck"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER
slide.notes_slide.notes_text_frame.text = "Speaker notes"

if visualizations:
    viz = visualizations[0]
    fields = [c['field'] if isinstance(c, dict) else str(c) for c in viz['columns']]
    rows = viz['rows']
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = [str(r.get(fields[0], '')) for r in rows]
    cd.add_series(fields[1], [float(r.get(fields[1]) or 0) for r in rows])
    chart = s2.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(2), Inches(11), Inches(5), cd).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.plots[0].has_data_labels = True
    chart.font.color.rgb = RGBColor(220, 220, 220)
    chart.category_axis.tick_labels.font.size = Pt(12)
    table = s2.shapes.add_table(len(rows) + 1, len(fields), Inches(1), Inches(0.2), Inches(6), Inches(1.5)).table
    for j, f in enumerate(fields):
        table.cell(0, j).text = f
    for i, r in enumerate(rows, start=1):
        for j, f in enumerate(fields):
            table.cell(i, j).text = str(r.get(f, ''))
    s2.shapes.title  # may be None on a blank layout; access must still validate

prs.save(_pptx_output_path)
"""

# 1x1 transparent PNG.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000100e221bc330000000049454e44ae426082"
)


def test_prompt_reference_deck_passes_validation_and_builds(tmp_path):
    out = tmp_path / f"{uuid.uuid4().hex}.pptx"
    validate_pptx_code(PROMPT_REFERENCE_DECK)
    PptxCodeExecutor().execute_pptx_code(
        code=PROMPT_REFERENCE_DECK,
        visualizations=[{
            "title": "Revenue by genre",
            "columns": [{"field": "Genre", "headerName": "Genre"}, {"field": "Revenue", "headerName": "Revenue"}],
            "rows": [{"Genre": "Rock", "Revenue": 826.65}, {"Genre": "Latin", "Revenue": 382.14}],
        }],
        report={"id": "r", "title": "Music Store revenue", "theme": None},
        output_path=out,
        images={"img-1": _PNG},
    )
    assert out.exists() and out.stat().st_size > 0


def test_ordinary_deck_code_still_generates(tmp_path):
    out = tmp_path / f"{uuid.uuid4().hex}.pptx"
    code = DECK_PREAMBLE + (
        "tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))\n"
        "tb.text_frame.text = type(prs).__name__\n"
        "tb.text_frame.paragraphs[0].font.size = Pt(24)\n"
        "prs.save(_pptx_output_path)\n"
    )
    _run(code, out)
    assert out.exists() and out.stat().st_size > 0
